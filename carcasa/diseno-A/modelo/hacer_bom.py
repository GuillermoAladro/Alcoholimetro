"""Lista de piezas (xlsx) de los 3 diseños: piezas a fabricar y compras, con fórmulas para totales."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
sys.path.insert(0, "/home/claude/nuevo")
import diseno

SP = "/tmp/claude-0/-home-claude/659a7a58-c69d-5ef4-bd94-d0aa972c769f/scratchpad/"
R = json.load(open(SP + "export_res.json"))
OUT = sys.argv[1]

# Compras que no están dibujadas en 3D (o cuya cantidad no sale del modelo): ref -> (uds por máquina, nota)
FIJAS = {
    "E13": (1, "Un rollo da para varias máquinas (unos 40 cm por máquina)"),
    "E14": (2, "Dos T: toma de presión y toma de alcohol"),
    "E15": (2, "Dos reductores del mismo pack (6→4 y 4→6) con un tramo de tubo de 4 mm entre ellos: es la barrera que crea la presión. El tubo de 4 mm no está en la lista"),
    "E17": (2, "Uno para la pantalla y otro para el Solo"),
    "E18": (1, "Se corta a 15 cm o menos"),
    "E19": (1, "UART y 5 V del sensor"),
    "E20": (1, "Fusibles de los 4 ramales: 3 A (Pi, Solo, pantalla) y 2 A (sensor)"),
    "E21": (1, ""), "E22": (1, "Cable de red del bar a la máquina"),
    "H08": (4, "Separadores de la cámara (dibujados con la cámara)"),
    "H11": (1, "Se cortan los frontales de policarbonato (pieza V01)"),
    "H13": (1, "Perfil de goma en los cantos del hueco del cajón de pajitas"),
    "H14": (18, "18 tirafondos a la pared (no dibujados), por los agujeros Ø6,5 de la trasera; se aprietan desde dentro con llave de vaso de 10"),
    "H15": (18, "Una por tirafondo"),
    "H16": (18, "Uno por tirafondo; broca de 8 mm y 60 mm de profundidad"),
}
# Material de montaje de la lista v7 (componentes-electronica.md), no dibujado
MONTAJE = [
    ("M01", "Fusibles cerámicos 5×20 T2,5A (para la C14)", "B0956XYTN6", 9.99, 1, 10, "Cambiar el de 10 A que trae la C14"),
    ("M02", "Cable 2×0,75 mm² rojo/negro, 10 m", "B0D12QPF4Y", 14.10, 1, None, "Troncal de 5 V y ramas"),
    ("M03", "Manguera H05VV-F 3G1,5 mm², 5 m", "B07BDHJGKG", 27.74, 1, None, "230 V y tierra (marrón, azul y verde-amarillo)"),
    ("M04", "Terminales de ojal M4 (25 uds.)", "B0CVW7Y2P4", 5.52, 6, 25, "Espárragos de tierra"),
    ("M05", "Arandelas dentadas M4 (200 uds.)", "B0B77J5FZS", 7.99, 4, 200, "Bajo cada ojal de tierra"),
    ("M06", "Tornillería M3–M6 inox (surtido)", "B0BGRY2JGX", 11.99, 1, None, "M3 fuente y bornes, M4 tuercas"),
    ("M07", "Terminales faston aislados (kit)", "B0B4H54KPS", 10.99, 1, None, "C14 y conector rápido del puente de tierra"),
    ("M08", "Cables Dupont (120 uds.)", "B01NGTXASZ", 9.99, 1, 120, "Pines de la Pi"),
    ("M09", "Crimpadora + punteras (herramienta)", "B0FQP6KDL4", 15.98, 1, None, "Herramienta: solo la primera vez"),
]
PACK = {"E15": 4, "H17": 20, "H14": 30, "H15": 25, "H16": 100, "E17": 1, "H01": 4, "H06": 20, "H07": 20, "H09": 350, "H10": 260, "E12": 2, "E09": 2, "E18": 6, "E08": 25}

AR = "Arial"
F = Font(name=AR, size=10)
FB = Font(name=AR, size=10, bold=True)
FH = Font(name=AR, size=10, bold=True, color="FFFFFF")
FT = Font(name=AR, size=14, bold=True)
FIN = Font(name=AR, size=10, color="0000FF")
FILL_H = PatternFill("solid", fgColor="2F3B46")
FILL_IN = PatternFill("solid", fgColor="FFF2B3")
FILL_TOT = PatternFill("solid", fgColor="E6EBEF")
thin = Side(style="thin", color="C6CFD6")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def cab(ws, fila, cols, anchos):
    for i, (c, w) in enumerate(zip(cols, anchos), 1):
        x = ws.cell(row=fila, column=i, value=c)
        x.font, x.fill, x.alignment, x.border = FH, FILL_H, Alignment(wrap_text=True, vertical="center"), BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[fila].height = 30


wb = Workbook()
res = wb.active
res.title = "Resumen"
hojas = {}
for L in sorted(R):
    d = R[L]
    # ---------------- piezas a fabricar
    ws = wb.create_sheet(f"{L} Fabricar")
    ws["A1"] = f"Diseño {L} · {d['nombre']} — piezas a fabricar"; ws["A1"].font = FT
    ws["A2"] = "Cada fila es una forma geométrica: una chapa plana (láser) o la placa de policarbonato. Archivos dentro de la carpeta CAD del diseño."
    ws["A2"].font = F
    cols = ["Nº pieza", "Pieza", "Grupo", "Material", "Proceso", "Largo (mm)", "Ancho (mm)", "Espesor (mm)", "Masa (g)", "Archivo DXF", "Archivo STEP", "Notas"]
    cab(ws, 4, cols, [10, 34, 20, 22, 18, 11, 11, 11, 10, 24, 26, 60])
    r0 = 5
    for i, f in enumerate(d["fabricadas"]):
        r = r0 + i
        vals = [f["id"], f["nombre"], f["grupo"], f["material"], f["proceso"], f["largo"], f["ancho"], f["espesor"], f["masa_g"], f["dxf"], f["step"], f["notas"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v); c.font = F; c.border = BOX; c.alignment = WRAP
        for j in (6, 7, 8):
            ws.cell(row=r, column=j).number_format = "0.0"
        ws.cell(row=r, column=9).number_format = "#,##0"
    r1 = r0 + len(d["fabricadas"]) - 1
    rt = r1 + 1
    ws.cell(row=rt, column=1, value="Total").font = FB
    ws.cell(row=rt, column=2, value=f"=COUNTA(A{r0}:A{r1})&\" piezas\"").font = FB
    ws.cell(row=rt, column=9, value=f"=SUM(I{r0}:I{r1})").font = FB
    ws.cell(row=rt, column=9).number_format = "#,##0"
    for j in range(1, 13):
        ws.cell(row=rt, column=j).fill = FILL_TOT
    ws.freeze_panes = "B5"
    hojas[(L, "fab")] = (ws.title, r0, r1, rt)

    # ---------------- compras
    wc = wb.create_sheet(f"{L} Comprar")
    wc["A1"] = f"Diseño {L} · {d['nombre']} — compras (Amazon.es)"; wc["A1"].font = FT
    wc["A2"] = "Celdas amarillas: paquetes a pedir (revisa cuántas unidades trae cada anuncio). Precios del anuncio consultados el 19–21/10/2026."
    wc["A2"].font = F
    cols = ["Ref.", "Artículo", "ASIN", "Enlace", "Uds. por máquina", "Uds. por anuncio", "Paquetes a pedir", "Precio anuncio (€)", "Coste (€)", "Números en el 3D", "Notas"]
    cab(wc, 4, cols, [7, 40, 13, 30, 11, 11, 11, 12, 11, 34, 52])
    filas = []
    modelo = {c["ref"]: c for c in d["compras"]}
    refs = sorted(set(list(modelo.keys()) + list(FIJAS.keys())))
    for ref in refs:
        n, asin, precio = diseno.COMPRAS[ref]
        if ref in modelo:
            uds = modelo[ref]["uds"]
            ids = ", ".join(x.split("-", 1)[1] for x in modelo[ref]["ids"])
        else:
            uds, ids = 0, "(no dibujado)"
        nota = ""
        if ref in FIJAS:
            uds = max(uds, FIJAS[ref][0]) if ref != "E17" else FIJAS[ref][0]
            nota = FIJAS[ref][1]
        if uds == 0:
            continue
        if ref == "H12":
            nota = "Sin anuncio fijado: buscar «tuerca ciega M4 inox» (DIN 1587)"
        if ref == "H17":
            nota = "Hace falta una punta Torx T20 de seguridad (con agujero); suele venir en los juegos de puntas"
        if ref == "H03":
            nota = "La leva recta se corta a 20,5 mm (ver documento)"
        if ref == "E16":
            nota = "Si el cable que trae la pantalla ya es acodado, este sobra"
        filas.append([ref, n, asin, f"https://www.amazon.es/dp/{asin}" if asin else "", uds, PACK.get(ref), None, precio, None, ids, nota])
    for m_ in MONTAJE:
        ref, n, asin, precio, uds, pack, nota = m_
        filas.append([ref, n, asin, f"https://www.amazon.es/dp/{asin}", uds, pack, None, precio, None, "(material de montaje)", nota])
    r0 = 5
    for i, f in enumerate(filas):
        r = r0 + i
        for j, v in enumerate(f, 1):
            c = wc.cell(row=r, column=j, value=v); c.font = F; c.border = BOX; c.alignment = WRAP
        if f[3]:
            wc.cell(row=r, column=4).hyperlink = f[3]
            wc.cell(row=r, column=4).font = Font(name=AR, size=10, color="1F5FBF", underline="single")
        # paquetes: ROUNDUP(uds/pack) si se conoce el pack; si no, 1 (a revisar)
        g = wc.cell(row=r, column=7, value=(f"=ROUNDUP(E{r}/F{r},0)" if f[5] else 1))
        g.font = FIN if not f[5] else F
        g.fill = FILL_IN
        wc.cell(row=r, column=8).number_format = "#,##0.00"
        wc.cell(row=r, column=9, value=f"=G{r}*H{r}").number_format = "#,##0.00"
        if f[5] is None:
            wc.cell(row=r, column=6).fill = FILL_IN
    r1 = r0 + len(filas) - 1
    rt = r1 + 1
    wc.cell(row=rt, column=1, value="Total").font = FB
    wc.cell(row=rt, column=9, value=f"=SUM(I{r0}:I{r1})").font = FB
    wc.cell(row=rt, column=9).number_format = "#,##0.00"
    for j in range(1, 12):
        wc.cell(row=rt, column=j).fill = FILL_TOT
    rm = rt + 1
    wc.cell(row=rm, column=1, value="Sin herramienta").font = F
    wc.cell(row=rm, column=9, value=f"=I{rt}-SUMIF(A{r0}:A{r1},\"M09\",I{r0}:I{r1})").number_format = "#,##0.00"
    wc.freeze_panes = "C5"
    hojas[(L, "com")] = (wc.title, r0, r1, rt)

# ---------------- resumen
res["A1"] = "Alcoholímetro · cajón de componentes · diseño " + ", ".join(sorted(R)); res["A1"].font = FT
res["A2"] = "Medidas exteriores en mm. Totales enlazados a las hojas de cada diseño (fórmulas)."; res["A2"].font = F
cols = ["Diseño", "Nombre", "Ancho (mm)", "Fondo (mm)", "Alto (mm)", "Tapa", "Piezas a fabricar", "Chapa y PC (kg)", "Compras primer pedido (€)",
        "Compras sin herramienta (€)", "Choques y barridos", "Idea"]
cab(res, 4, cols, [8, 14, 10, 10, 10, 34, 11, 11, 14, 14, 26, 70])
for i, L in enumerate(sorted(R)):
    r = 5 + i
    d = R[L]
    tf, f0, f1, ft = hojas[(L, "fab")]
    tc, c0, c1, ct = hojas[(L, "com")]
    tapa = "Abatible hacia la izquierda, 1 cerradura" if (d["tapa"] == "lateral" and d["bisagra"] == "izq") else \
           "Abatible hacia la derecha, 1 cerradura" if d["tapa"] == "lateral" else "Se quita (lengüetas detrás), 2 cerraduras"
    chk = sum(d["barridos"].values()) + d["choques"] + d["zonas"]
    vals = [L, d["nombre"], d["cotas"]["W"], d["cotas"]["D"], d["cotas"]["H"], tapa,
            f"=COUNTA('{tf}'!A{f0}:A{f1})", f"='{tf}'!I{ft}/1000", f"='{tc}'!I{ct}", f"='{tc}'!I{ct + 1}",
            ("0 choques · tapa, tapa interior, cajón y pajitas sin roces" if chk == 0 else f"{chk} avisos"), d["descripcion"]]
    for j, v in enumerate(vals, 1):
        c = res.cell(row=r, column=j, value=v); c.font = F; c.border = BOX; c.alignment = WRAP
    res.cell(row=r, column=8).number_format = "0.0"
    res.cell(row=r, column=9).number_format = "#,##0.00"
    res.cell(row=r, column=10).number_format = "#,##0.00"
    for j in (3, 4, 5):
        res.cell(row=r, column=j).number_format = "0"
res["A10"] = "Notas"; res["A10"].font = FB
notas = [
    "Hay dos policarbonatos de 5 mm: uno delante del panel frontal y otro detrás, entre el panel y el cajón de componentes. Todas las piezas de chapa son planas (acero DC01 1,5 mm, corte láser) y se unen por soldadura por puntos o remaches: no hay plegados.",
    "Los números de pieza son los mismos en el visor 3D, en los STEP (STEP_piezas/<nº>.step) y en los DXF (DXF_corte/<nº>.dxf).",
    "Las compras salen del modelo 3D (cuántas veces aparece cada pieza) más lo que no se dibuja (tubo, fusibles, cables, microSD...).",
    "Paquetes a pedir: si el anuncio no dice cuántas unidades trae, se pone 1 en amarillo para revisarlo.",
]
for i, t in enumerate(notas):
    c = res.cell(row=11 + i, column=1, value="• " + t); c.font = F
res.freeze_panes = "A5"
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("ok", OUT)
